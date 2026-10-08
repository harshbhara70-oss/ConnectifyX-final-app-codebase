import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set cell padding (in dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
    """Set clean subtle borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_flutter_backend_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Tokens
    COLOR_NAVY = RGBColor(30, 58, 138)       # Primary
    COLOR_TEAL = RGBColor(13, 148, 136)      # Secondary Accent
    COLOR_SLATE = RGBColor(31, 41, 55)       # Text
    COLOR_MUTED = RGBColor(75, 85, 99)       # Subdued Text
    HEADER_BG_HEX = "1E3A8A"
    ALT_ROW_HEX = "F8FAFC"

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("CONNECTIFYX")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("Flutter Mobile Application & Backend Engineering Handover Report")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_TEAL

    # Header Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    set_table_borders(meta_table, color="E2E8F0", sz="4")

    col_widths = [Inches(2.2), Inches(4.6)]
    metadata = [
        ("Platform Name:", "CONNECTIFYX Mobile App & Cloud Backend"),
        ("Architecture Scope:", "Flutter Native Client (c2c_android) + Express 5 Node Backend + Gemini AI"),
        ("Version & Release:", "v2.4.0-Production Ready"),
        ("Target Recipient:", "Incoming Mobile & Backend Engineering Interns")
    ]

    for r_idx, (lbl, val) in enumerate(metadata):
        row = meta_table.rows[r_idx]
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        
        # Label
        c_lbl = row.cells[0]
        set_cell_background(c_lbl, "F1F5F9")
        p_l = c_lbl.paragraphs[0]
        p_l.paragraph_format.space_before = Pt(0)
        p_l.paragraph_format.space_after = Pt(0)
        r_l = p_l.add_run(lbl)
        r_l.font.name = "Arial"
        r_l.font.size = Pt(9.5)
        r_l.font.bold = True
        r_l.font.color.rgb = COLOR_NAVY

        # Value
        c_val = row.cells[1]
        set_cell_background(c_val, "FFFFFF")
        p_v = c_val.paragraphs[0]
        p_v.paragraph_format.space_before = Pt(0)
        p_v.paragraph_format.space_after = Pt(0)
        r_v = p_v.add_run(val)
        r_v.font.name = "Arial"
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = COLOR_SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Helper Functions
    def add_h(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        if level == 1:
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(14.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = COLOR_TEAL
        elif level == 3:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_SLATE
        return p

    def add_p(text, bold_prefix=None, bullet=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        if bullet:
            p.paragraph_format.left_indent = Inches(0.25)
            dot = p.add_run("•  ")
            dot.font.name = "Arial"
            dot.font.size = Pt(10)
            dot.font.color.rgb = COLOR_TEAL
            dot.font.bold = True

        if bold_prefix:
            bp_run = p.add_run(bold_prefix)
            bp_run.font.name = "Arial"
            bp_run.font.size = Pt(10)
            bp_run.font.bold = True
            bp_run.font.color.rgb = COLOR_SLATE

        body_run = p.add_run(text)
        body_run.font.name = "Arial"
        body_run.font.size = Pt(10)
        body_run.font.color.rgb = COLOR_SLATE
        return p

    def add_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, color="CBD5E1", sz="4")

        # Header Row
        hdr_row = table.rows[0]
        for idx, h_text in enumerate(headers):
            cell = hdr_row.cells[idx]
            if col_widths and idx < len(col_widths):
                cell.width = col_widths[idx]
            set_cell_background(cell, HEADER_BG_HEX)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(h_text)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_vals in enumerate(rows_data):
            row = table.rows[r_idx + 1]
            bg_color = ALT_ROW_HEX if r_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_vals):
                cell = row.cells[c_idx]
                if col_widths and c_idx < len(col_widths):
                    cell.width = col_widths[c_idx]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(str(val))
                run.font.name = "Arial"
                run.font.size = Pt(9)
                run.font.color.rgb = COLOR_SLATE

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return table

    # -------------------------------------------------------------
    # 1. Executive Summary: Mobile & Backend Ecosystem
    # -------------------------------------------------------------
    add_h("1. Executive Summary & Mobile-Backend Architecture", level=1)
    add_p("The CONNECTIFYX mobile ecosystem consists of an enterprise Flutter cross-platform mobile application communicating with a high-throughput Express.js 5 REST API backend, powered by MongoDB Atlas and Google Gemini 2.5 Flash Artificial Intelligence.")
    add_p("This report is exclusively compiled for incoming mobile engineers and backend developers to understand the architecture, data models, state management, REST endpoints, AI pipelines, and operational setup.")

    # -------------------------------------------------------------
    # 2. Flutter Mobile Application (c2c_android)
    # -------------------------------------------------------------
    add_h("2. Flutter Mobile Application Architecture (c2c_android)", level=1)
    add_p("The mobile client is built with Flutter 3.11+ using the BLoC (Business Logic Component) pattern for robust, predictable, reactive state management across all five user roles.")

    add_h("2.1 Mobile Tech Stack & Dependencies (pubspec.yaml)", level=2)
    flut_headers = ["Package Name", "Version", "Architectural Role", "Benefit & Usage"]
    flut_data = [
        ["flutter_bloc", "^8.1.6", "State Management", "Decouples business logic from UI widgets with events & states"],
        ["dio", "^5.7.0", "HTTP Networking", "Intercepts API requests, auto-injects Bearer JWT, handles refresh/retry"],
        ["go_router", "^14.6.2", "Declarative Navigation", "URL-based path routing with role-based auth redirect guards"],
        ["flutter_secure_storage", "^9.2.2", "Encrypted Storage", "Persists JWT tokens securely on Android KeyStore & iOS Keychain"],
        ["shared_preferences", "^2.3.5", "Local Cache", "Caches UI theme mode (light/dark) and user preferences"],
        ["fl_chart", "^0.69.0", "Data Visualization", "Renders animated student readiness radial bars & salary graphs"],
        ["pdf & printing", "^3.12.0 / ^5.14.3", "Document Export", "Generates and prints ATS-compliant student resume PDFs on-device"],
        ["shimmer & lottie", "^3.0.0 / ^3.1.2", "User Feedback", "Skeleton loading placeholders and animated micro-interactions"],
        ["flutter_displaymode", "^0.6.0", "Performance", "Unlocks 90Hz/120Hz high refresh-rate displays for fluid 60fps+ UI"],
        ["google_fonts", "^6.2.0", "Typography", "Consistent typography across Android and iOS rendering engines"]
    ]
    add_table(flut_headers, flut_data, [Inches(1.8), Inches(1.1), Inches(1.8), Inches(2.1)])

    add_h("2.2 Mobile Directory Structure & Layering (lib/)", level=2)
    add_p("The Flutter application enforces a strict feature-first clean architecture:")
    add_p(" Contains singleton DioClient, ApiInterceptors, AppRouter, AppTheme, and CustomBottomNav.", bold_prefix="• lib/core/: ")
    add_p(" Modularized into 6 role-specific sub-packages containing their own BLoC, Models, Repositories, Pages, and Widgets:", bold_prefix="• lib/features/: ")
    add_p(" Authentication, Token Guard, Role Selection, Login/Register pages.", bold_prefix="    1. features/auth/: ")
    add_p(" Student Cockpit, AI Resume, Ask AI, Mock Interview, Projects, Certificates, Resume PDF Generator.", bold_prefix="    2. features/student/: ")
    add_p(" Recruiter Dashboard, Job Postings, Candidate Pipeline Kanban, Direct Messages.", bold_prefix="    3. features/recruiter/: ")
    add_p(" TPO Dashboard, Student Records, Departmental Analytics, At-Risk Early Warning, Drive Management.", bold_prefix="    4. features/college/: ")
    add_p(" 360° Employability Matrix, Session Timeline, Feedback Scoring, Study Resources, Leaderboard.", bold_prefix="    5. features/mentor/: ")
    add_p(" Super Admin Dashboard, User Governance, Verification Queue, Global Broadcasts, Analytics.", bold_prefix="    6. features/admin/: ")

    add_h("2.3 Mobile Screens & Feature Breakdown", level=2)
    
    add_h("🎓 Student Mobile Features:", level=3)
    stud_screens = [
        ("Student Dashboard (student_dashboard_page.dart):", " Real-time radial readiness gauge, daily activity streaks, upcoming drive alerts, and quick action cards."),
        ("AI Resume Studio (ai_resume_page.dart):", " Upload/type resume text to receive instant ATS compatibility score, missing keyword highlights, and AI bullet rewrites."),
        ("Ask AI Career Coach (ask_ai_page.dart):", " Interactive conversational chat assistant powered by Gemini 2.5 Flash for career advice and salary negotiation tips."),
        ("Placement Mock Interview (student_hiring_process_page.dart):", " Interactive mock technical/HR rounds with instant AI answer grading and feedback."),
        ("Capstone Project Hub (student_project_list_page.dart / student_applied_projects_page.dart):", " Browse and submit capstones with GitHub repository links and review status."),
        ("Native Resume PDF Exporter (resume_pdf_generator.dart):", " Generates printable, ATS-optimized PDF resumes directly on the mobile device.")
    ]
    for stitle, sdesc in stud_screens:
        add_p(sdesc, bold_prefix=stitle, bullet=True)

    add_h("🏢 Recruiter Mobile Features:", level=3)
    rec_screens = [
        ("Recruiter Dashboard (recruiter_dashboard_page.dart):", " Tracks active campus jobs, total applications, and round-wise conversion funnel."),
        ("Job Posting Creator (post_job_page.dart / my_job_postings_page.dart):", " Publish and manage drives specifying package (LPA), CGPA cutoff, and branch criteria."),
        ("Pipeline Kanban Manager (pipeline_management_page.dart):", " Move candidates between stages (Applied -> Shortlisted -> Technical -> HR -> Offered)."),
        ("Recruiter Messaging (recruiter_messages_page.dart):", " Direct chat interface with shortlisted applicants.")
    ]
    for rtitle, rdesc in rec_screens:
        add_p(rdesc, bold_prefix=rtitle, bullet=True)

    add_h("🏛️ College / TPO Mobile Features:", level=3)
    col_screens = [
        ("College Placement Telemetry (college_dashboard_page.dart):", " Departmental placement statistics, highest package, and active recruiter counts."),
        ("Student Cohort Records (student_records_page.dart / college_student_detail_page.dart):", " Comprehensive student directory with multi-attribute filtering."),
        ("At-Risk Early Warning (at_risk_students_page.dart):", " Flags unplaced students with low test scores and missing resumes for proactive coaching."),
        ("Departmental Comparisons & Batch Grouping:", " Compare CSE vs IT vs ECE placement velocity and manage custom student batches.")
    ]
    for ctitle, cdesc in col_screens:
        add_p(cdesc, bold_prefix=ctitle, bullet=True)

    add_h("👨‍🏫 Mentor Mobile Features:", level=3)
    men_screens = [
        ("360° Employability Matrix (mentor_employability_matrix_page.dart):", " Multi-axial scoring interface for student technical and communication skills."),
        ("Session Timeline & Scheduling (mentor_session_timeline_page.dart):", " View booked 1-on-1 sessions, generate video call links, and track attendance."),
        ("Study Materials & Resources (mentor_study_material_page.dart):", " Upload coding cheat-sheets, interview questions, and placement roadmaps.")
    ]
    for mtitle, mdesc in men_screens:
        add_p(mdesc, bold_prefix=mtitle, bullet=True)

    add_h("🛡️ Super Admin Mobile Features:", level=3)
    adm_screens = [
        ("Platform Governance (admin_dashboard_page.dart / admin_user_management_page.dart):", " Master user list across all roles with instant suspension/activation controls."),
        ("Verification Queue (admin_verification_queue_page.dart):", " Institutional KYB verification of AISHE college credentials and recruiter domains."),
        ("Global Broadcast Control (admin_broadcast_control_page.dart):", " Push platform-wide alerts and placement announcements to all mobile users.")
    ]
    for atitle, adesc in adm_screens:
        add_p(adesc, bold_prefix=atitle, bullet=True)

    # -------------------------------------------------------------
    # 3. Backend Server Architecture (backend/)
    # -------------------------------------------------------------
    add_h("3. Backend Server Architecture (backend/)", level=1)
    add_p("The backend is built with Express.js 5 on Node.js (v18+) using modern ECMAScript Modules (ESM). It communicates with MongoDB Atlas via Mongoose 9.x.")

    add_h("3.1 Backend Middleware Pipeline", level=2)
    add_p(" Configures allowed origins and headers for seamless mobile and client access.", bold_prefix="• CORS Middleware (cors): ")
    add_p(" Parses incoming JSON payloads and URL-encoded bodies up to 10MB.", bold_prefix="• Body Parsers (express.json): ")
    add_p(" Extracts and validates JWT tokens from the Authorization: Bearer <token> header, verifying user role signatures.", bold_prefix="• Auth Guards (authMiddleware, studentAuthMiddleware, adminAuth): ")
    add_p(" Enforces role-based permissions (e.g. only recruiters can post jobs; only admins can verify institutions).", bold_prefix="• Role Authorization (roleMiddleware): ")
    add_p(" Intercepts uncaught exceptions and MongoDB validation errors, returning structured JSON responses.", bold_prefix="• Global Error Handler (errorHandler.js): ")

    # -------------------------------------------------------------
    # 4. Google Gemini GenAI Microservices
    # -------------------------------------------------------------
    add_h("4. Google Gemini 2.5 Flash Artificial Intelligence Microservices", level=1)
    add_p("The backend integrates the official @google/genai SDK in backend/src/services/aiServices.js using the gemini-2.5-flash model with responseMimeType: 'application/json' for structured responses:")

    ai_headers = ["AI Service Function", "Input Context", "Structured JSON Output", "Core Purpose"]
    ai_data = [
        [
            "analyzeResume(resumeText)",
            "Raw resume text extracted from user input / PDF",
            '{"score": 82, "title": "...", "keywords_found": [...], "keywords_missing": [...], "tip": "..."}',
            "Evaluates ATS compatibility, detects missing keywords, and highlights missing skills."
        ],
        [
            "generateStudyPlan(studentContext)",
            "Student branch, CGPA, target role, weak areas",
            '[{"day": "Monday", "tasks": ["Task 1", "Task 2", "Task 3"]}, ...]',
            "Generates customized 5-day intensive placement preparation schedule."
        ],
        [
            "analyzeSkillGap(studentContext, role)",
            "Student profile skills vs. target job role",
            '{"match_score": 78, "matched_skills": [...], "missing_skills": [...], "suggested_modules": [...]}',
            "Identifies technical gaps and recommends specific learning modules."
        ],
        [
            "careerCoach(question, context)",
            "Student question + academic context",
            '{"answer": "Concise, actionable career advisory."}',
            "Natural language AI placement mentor answering queries 24/7."
        ],
        [
            "generateHiringInterview(context)",
            "Target role and company profile",
            '{"title": "...", "rounds": [{"name": "Technical Round", "questions": [...]}]}',
            "Dynamically generates realistic interview rounds with technical/HR questions."
        ],
        [
            "evaluateHiringAnswer(context)",
            "Interview question + student response",
            '{"score": 85, "strengths": [...], "improvements": [...], "feedback": "..."}',
            "Scores candidate answer quality, clarity, and technical correctness."
        ]
    ]
    add_table(ai_headers, ai_data, [Inches(1.8), Inches(1.5), Inches(2.2), Inches(1.7)])

    # -------------------------------------------------------------
    # 5. Database Schema & Data Models
    # -------------------------------------------------------------
    add_h("5. Database Collections & Schema Architecture", level=1)
    add_p("The database consists of 11 core Mongoose models located in backend/src/models/:")

    model_headers = ["Model Name", "Schema Fields Summary", "Relations & Foreign Keys", "Storage Role"]
    model_data = [
        ["student.js", "name, email, password, college, branch, semester, gpa, skills[], placementStatus, metrics{}", "Ref by Application, MentorEvaluation", "Student identity, academic history & readiness scores"],
        ["recruiter.js & company.js", "companyName, website, industry, contactEmail, isVerified, logo, location", "Parent of Project/JobDrive", "Recruiter organization and verification status"],
        ["project.js", "title, companyId, role, packageLPA, eligibility{minGpa, allowedBranches[]}, deadline", "Ref -> Company; Ref by Application", "Active campus job drives and recruitment projects"],
        ["application.js", "studentId, jobId, status ('Applied'|'Shortlisted'|'Technical'|'HR'|'Selected'), interviewDate", "Ref -> Student, Ref -> Project", "Full lifecycle application stage progression tracking"],
        ["college.js", "collegeName, aisheCode, accreditation, verificationStatus, contactEmail, departments[]", "Ref by Student", "Institutional accreditation & TPO records"],
        ["mentorProfile.js", "name, email, company, role, expertise[], bio, rating, totalMentees", "Parent of MentorSession", "Industry mentor profile & domain specializations"],
        ["mentorSession.js", "mentorId, studentId, scheduledTime, topic, meetingLink, status", "Ref -> MentorProfile, Ref -> Student", "1-on-1 mentorship appointments & mock interview logs"],
        ["mentorEvaluation.js", "mentorId, studentId, technicalRating, softSkillRating, communicationRating, notes[]", "Ref -> Mentor, Ref -> Student", "360° Employability rating entries and mentor feedback"],
        ["recruiterMessage.js", "senderId, recipientId, messageText, timestamp, readStatus", "Ref -> Recruiter, Ref -> Student", "Real-time communication logs between recruiters & students"],
        ["admin.js", "email, password, role ('superadmin'), permissions[], lastLogin", "System Auth", "Root platform administrators and security oversight"]
    ]
    add_table(model_headers, model_data, [Inches(1.4), Inches(2.3), Inches(1.7), Inches(1.8)])

    # -------------------------------------------------------------
    # 6. REST API Endpoints Specification
    # -------------------------------------------------------------
    add_h("6. REST API Endpoints Specification", level=1)
    add_p("All mobile app interactions are routed to http://localhost:5000/api via DioClient:")

    api_headers = ["Method", "Endpoint Path", "Controller Function", "Auth Guard", "Action Performed"]
    api_data = [
        ["POST", "/api/auth/register", "studentController.register", "Public", "Register new student with encrypted password"],
        ["POST", "/api/auth/login", "studentController.login", "Public", "Authenticate student and issue JWT"],
        ["GET", "/api/student/profile", "studentController.getProfile", "Student JWT", "Fetch student profile and metrics"],
        ["PUT", "/api/student/profile", "studentController.updateProfile", "Student JWT", "Update skills, bio, and academic marks"],
        ["GET", "/api/student/applications", "studentController.getApplications", "Student JWT", "Get list of all submitted drive applications"],
        ["POST", "/api/student/apply/:jobId", "studentController.applyJob", "Student JWT", "Apply for an active campus drive"],
        ["POST", "/api/ai/resume-analyzer", "aiController.analyzeResume", "Student JWT", "Parse resume & return ATS analysis"],
        ["POST", "/api/ai/study-plan", "aiController.getStudyPlan", "Student JWT", "Generate 5-day AI preparation roadmap"],
        ["POST", "/api/ai/skill-gap", "aiController.getSkillGap", "Student JWT", "Calculate role match % and missing skills"],
        ["POST", "/api/ai/career-coach", "aiController.coachChat", "Student JWT", "Conversational AI placement coaching"],
        ["POST", "/api/ai/interview-gen", "aiController.generateInterview", "Student JWT", "Generate mock technical interview questions"],
        ["POST", "/api/ai/evaluate-answer", "aiController.evaluateAnswer", "Student JWT", "Grade student response to interview question"],
        ["POST", "/api/recruiter/jobs", "recruiterController.createJob", "Recruiter JWT", "Publish new campus job opening"],
        ["GET", "/api/recruiter/jobs", "recruiterController.getMyJobs", "Recruiter JWT", "Fetch jobs created by the recruiter's company"],
        ["GET", "/api/recruiter/pipeline/:jobId", "recruiterController.getPipeline", "Recruiter JWT", "Retrieve candidate pool grouped by hiring round"],
        ["PATCH", "/api/recruiter/application/:id/status", "recruiterController.updateStatus", "Recruiter JWT", "Advance applicant stage (e.g. to Technical, Hired)"],
        ["POST", "/api/mentor/sessions", "mentorController.createSession", "Mentor JWT", "Schedule 1-on-1 mentorship session"],
        ["POST", "/api/mentor/evaluation", "mentorController.submitEvaluation", "Mentor JWT", "Submit 360° student matrix score"],
        ["POST", "/api/admin/login", "adminController.login", "Public", "Super Admin JWT authentication"],
        ["GET", "/api/admin/users", "adminController.getAllUsers", "Admin JWT", "Fetch all users across all 5 portals"],
        ["PATCH", "/api/admin/verify-entity", "adminController.verifyEntity", "Admin JWT", "Approve/Reject College or Recruiter registration"]
    ]
    add_table(api_headers, api_data, [Inches(0.8), Inches(2.2), Inches(1.6), Inches(1.0), Inches(1.6)])

    # -------------------------------------------------------------
    # 7. Step-by-Step Developer Setup & Run Guide
    # -------------------------------------------------------------
    add_h("7. Step-by-Step Developer Setup & Run Guide", level=1)
    add_p("To run both the Backend API and Flutter Mobile App on a new development workstation:")

    add_h("Step 1: Configure Backend Environment (.env)", level=2)
    add_p("Ensure .env exists in the project root with the following configuration:")
    env_str = (
        "MONGO_URI=mongodb+srv://<username>:<password>@cluster0.mongodb.net/connectifyx?retryWrites=true&w=majority\n"
        "GEMINI_API_KEY=your_gemini_api_key_here\n"
        "GEMINI_MODEL=gemini-2.5-flash\n"
        "ADMIN_EMAIL=admin@connectifyx.com\n"
        "ADMIN_PASSWORD=your_admin_secure_password\n"
        "JWT_ADMIN_SECRET=your_super_secret_admin_jwt_key_2026\n"
        "PORT=5000"
    )
    p_e = doc.add_paragraph()
    p_e.paragraph_format.left_indent = Inches(0.25)
    r_e = p_e.add_run(env_str)
    r_e.font.name = "Consolas"
    r_e.font.size = Pt(8.5)
    r_e.font.color.rgb = COLOR_SLATE

    add_h("Step 2: Start the Express.js Backend Server", level=2)
    p_b = doc.add_paragraph()
    p_b.paragraph_format.left_indent = Inches(0.25)
    r_b = p_b.add_run("cd backend\nnpm install\nnpm run dev")
    r_b.font.name = "Consolas"
    r_b.font.size = Pt(9.5)
    r_b.font.bold = True
    r_b.font.color.rgb = COLOR_NAVY
    add_p("Verify console output: 'Server running on port 5000' & 'MongoDB Connected Successfully'.")

    add_h("Step 3: Run the Flutter Mobile Application (c2c_android)", level=2)
    p_f = doc.add_paragraph()
    p_f.paragraph_format.left_indent = Inches(0.25)
    r_f = p_f.add_run("cd c2c_android\nflutter pub get\nflutter run")
    r_f.font.name = "Consolas"
    r_f.font.size = Pt(9.5)
    r_f.font.bold = True
    r_f.font.color.rgb = COLOR_NAVY
    add_p("The Flutter app will compile and launch on your connected Android device, iOS simulator, or desktop client.")

    # -------------------------------------------------------------
    # 8. Critical Onboarding Notes & Roadmap
    # -------------------------------------------------------------
    add_h("8. Critical Onboarding Notes & Roadmap for Incoming Interns", level=1)
    
    add_h("8.1 Critical Technical Notes", level=2)
    add_p(" When testing on an Android Emulator, use 10.0.2.2:5000 instead of localhost:5000. When testing on a physical device, ensure the phone and computer are on the same Wi-Fi and use your local machine's IP (e.g., 192.168.1.X:5000 in lib/core/constants/api_endpoints.dart).", bold_prefix="1. Android Emulator Loopback IP:", bullet=True)
    add_p(" Tokens are stored encrypted in flutter_secure_storage. To test cold sign-ins, use the 'Logout' action in settings which clears the keychain.", bold_prefix="2. Secure Storage Cache:", bullet=True)
    add_p(" All Gemini AI responses must return strict valid JSON matching the schema in backend/src/services/aiServices.js.", bold_prefix="3. Gemini JSON Mode:", bullet=True)

    add_h("8.2 Mobile-Backend Roadmap & Upcoming Enhancements", level=2)
    roadmaps = [
        ("Phase 1 — WebSocket Push Notifications:", " Integrate Socket.io client in Flutter to listen for real-time application status updates and instant recruiter chat messages."),
        ("Phase 2 — Biometric Authentication:", " Add local_auth to enable fingerprint / FaceID unlock for students and recruiters on mobile."),
        ("Phase 3 — Offline Synced Mode:", " Implement Hive or Isar database caching for offline viewing of applied projects and saved study roadmaps."),
        ("Phase 4 — Real-time Video Mock Interviews:", " Integrate WebRTC live media streaming for real-time video interview practice with speech-to-text analysis.")
    ]
    for r_t, r_d in roadmaps:
        add_p(r_d, bold_prefix=r_t, bullet=True)

    # Signoff footer
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(16)
    p_foot.paragraph_format.space_after = Pt(0)
    r_foot = p_foot.add_run("Handover Complete & Verified for Incoming Mobile & Backend Developers — CONNECTIFYX")
    r_foot.font.name = "Arial"
    r_foot.font.size = Pt(10)
    r_foot.font.bold = True
    r_foot.font.italic = True
    r_foot.font.color.rgb = COLOR_NAVY

    # Output file
    output_path = "CONNECTIFYX_Flutter_App_and_Backend_Handover_Report.docx"
    doc.save(output_path)
    print(f"Generated: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    build_flutter_backend_report()
